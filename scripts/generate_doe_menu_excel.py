#!/usr/bin/env python3
"""Generate Smartsheet-ready Excel from UNM Day of Event menu PDF."""

import openpyxl
from openpyxl import load_workbook
from copy import copy

AUDIT_PATH = "/home/ubuntu/.cursor/projects/workspace/uploads/UNM_2026_Item_Setup__Audit__-_EEG_8c88.xlsx"
OUTPUT_PATH = "/workspace/UNM_DOE_Menu_2026_Smartsheet.xlsx"

# (excel_category, [(oms, menu_name, price, description), ...])
MENU_SECTIONS = [
    (
        "HOT APPETIZER",
        [
            (
                2100882,
                "Classic Chips, Chile Con Queso & Salsa Fresca",
                175,
                "Warm Chile con Queso, House-Made Salsa, House-Made Tortilla Chips",
            ),
            (
                2101228,
                "Traditional Nacho Platter",
                175,
                "Picadillo, Monterey Jack Cheese, Sour Cream, Avocado, Diced Tomatoes, Cilantro, Red Salsa, House-Made Tortilla Chips",
            ),
            (
                2102535,
                "The Chicken Tender & Sauce Zone",
                210,
                "Golden Brown Chicken Tenders, & Sauce Bar: New Mexico Bang Bang Sauce, Signature Barbecue Sauce, Buttermilk Ranch Dressing, Spicy Buffalo Sauce",
            ),
            (
                230,
                "Spicy Wings",
                220,
                "Traditional Spicy Buffalo Sauce, New Mexico Bang Bang Sauce, Buttermilk Ranch Dressing",
            ),
        ],
    ),
    (
        "SNACKS",
        [
            (
                14322,
                "Louie's Snack Attack",
                105,
                "M&M's Peanuts, Pretzels, Levy Snack Mix, Freshly Popped Popcorn",
            ),
            (
                3200010,
                "Salsa & Guacamole Sampler",
                135,
                "Salsa, Fresh Guacamole, House-Made Tortilla Chips",
            ),
            (13333, "Bottomless Freshly Popped Popcorn", 65, None),
            (164, "Snack Mix", 45, None),
            (163, "Pretzel Twists", 45, None),
            (125, "Dry-Roasted Peanuts", 45, None),
            (
                6651,
                "Gourmet Cookies & Brownies",
                170,
                "Assorted Gourmet Cookies & Decadent Brownies; Serves 10 guests",
            ),
        ],
    ),
    (
        "HOT DOGS/SAUSAGES",
        [
            (
                3000328,
                "Lobo Dogs",
                180,
                "Grilled Hot Dogs, Traditional Condiments; 2 DOGS PER GUEST",
            ),
        ],
    ),
    (
        "UPCHARGE",
        [
            (
                99216,
                "Upgrade your Lobo Dogs: Frito Pie Red Chile",
                40,
                "Add-on upgrade for Lobo Dogs",
            ),
        ],
    ),
    (
        "ENTREES",
        [
            (
                2303580,
                "Louie's Legendary Frito Pie",
                170,
                "Chile, Fritos Corn Chips, Shredded Lettuce, Tomatoes, Diced Onions, Shredded Cheese",
            ),
        ],
    ),
    (
        "BEV-SOFT DRINKS",
        [
            (904, "Pepsi", 45, None),
            (4700138, "Pepsi Zero Sugar", 45, None),
            (4700623, "Starry", 45, None),
            (918, "Mug Root Beer", 45, None),
            (4700098, "Pure Leaf Iced Tea Lemon 18.5oz", 65, None),
            (4700055, "Pure Leaf Sweet Tea 18.5oz", 65, None),
            (4700056, "Pure Leaf Unsweetened Black Tea 18.5oz", 65, None),
        ],
    ),
    (
        "BEV-WATER",
        [
            (958, "Aquafina Bottled Water", 40, None),
            (970, "Schweppes Club Soda", 25, None),
        ],
    ),
    (
        "BAR MIXERS",
        [
            (5000178, "Lemon Wedges", 20, None),
            (5000134, "Lime Wedges", 20, None),
            (5000180, "Orange Wedges", 20, None),
        ],
    ),
    (
        "BEER-DOMESTIC",
        [
            (704, "Bud Light", 60, None),
            (1730, "Michelob ULTRA", 60, None),
            (6001392, "Marble Double White", 85, None),
            (
                6001205,
                "Classic Bundle",
                85,
                "Bud Light, Michelob ULTRA, Estrella Jalisco; 85 PER 6-PACK",
            ),
            (
                6003196,
                "505 Craft Collection",
                85,
                "A Microbrew Bundle from Local Breweries: Santa Fe Brewing Co. 7K IPA, Teller Lobo Cherry Limeade, NÜTRL Vodka Seltzer; 85 PER 6-PACK",
            ),
            (
                None,
                "Southwest Select",
                85,
                "Modelo Especial, Michelob ULTRA; 85 PER 6-PACK",
            ),
        ],
    ),
    (
        "BEER-IMPORTED",
        [
            (77940, "Dos Equis Lager", 85, None),
        ],
    ),
    (
        "BEER-RTD",
        [
            (6002607, "NÜTRL Vodka Seltzer", 85, None),
        ],
    ),
    (
        "WINE-SPARKLING",
        [
            (3700127, "Gruet Brut", 110, "Sparkling"),
            (None, "Cavit", 110, "Sparkling"),
        ],
    ),
    (
        "WINE-WHITE",
        [
            (3800357, "Woodbridge Pinot Grigio", 60, None),
            (16090, "Chateau Ste. Michelle Riesling", 100, None),
            (13659, "Kim Crawford Sauvignon Blanc", 130, None),
            (12000, "Kendall-Jackson Chardonnay", 125, None),
            (11198, "Woodbridge Chardonnay", 60, None),
        ],
    ),
    (
        "WINE-RED",
        [
            (3900152, "Meiomi Pinot Noir", 115, None),
            (3900724, "Emmolo Merlot", 195, None),
            (11194, "Woodbridge Merlot", 60, None),
            (3901227, "Decoy by Duckhorn Cabernet Sauvignon", 190, None),
            (16029, "Louis Martini Cabernet Sauvignon", 100, None),
            (11189, "Woodbridge Cabernet Sauvignon", 60, None),
            (3900051, "The Prisoner Red Blend", 280, None),
        ],
    ),
]


def get_audit_row(ws_audit, oms):
    for row in range(2, ws_audit.max_row + 1):
        if ws_audit.cell(row=row, column=2).value == oms:
            return [ws_audit.cell(row=row, column=c).value for c in range(1, 26)]
    return None


def blank_row():
    return [None] * 25


def build_item_row(ws_audit, oms, menu_name, menu_price, description=None):
    base = get_audit_row(ws_audit, oms) if oms else blank_row()
    if base is None:
        base = blank_row()

    row = list(base)
    row[0] = None
    if oms:
        row[1] = oms
    row[2] = menu_name
    row[6] = menu_price
    row[7] = menu_price
    if description:
        row[17] = description
    return row


def main():
    wb_audit = load_workbook(AUDIT_PATH)
    ws_audit = wb_audit["UNM 2026 Item Setup (Audit)"]

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "UNM 2026 Item Setup (Audit)"

    for col in range(1, 26):
        src = ws_audit.cell(row=1, column=col)
        dst = ws_out.cell(row=1, column=col, value=src.value)
        if src.has_style:
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.alignment = copy(src.alignment)

    for excel_cat, items in MENU_SECTIONS:
        ws_out.append([excel_cat] + [None] * 24)
        for oms, name, price, desc in items:
            ws_out.append(build_item_row(ws_audit, oms, name, price, desc))

    wb_out.create_sheet("Comments")
    wb_out.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} with {ws_out.max_row - 1} data rows")


if __name__ == "__main__":
    main()
