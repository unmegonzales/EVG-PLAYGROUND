#!/usr/bin/env python3
"""Build a portrait 8.5x11 PowerPoint handout of The Pit 60th cup concepts."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "The-Pit-60-Souvenir-Cup-Executive-Review.pptx"

INK = RGBColor(0x11, 0x11, 0x11)
CHERRY = RGBColor(0xBA, 0x0C, 0x2F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xF5, 0xF0, 0xE8)
SILVER = RGBColor(0xA7, 0xA8, 0xAA)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
WARM = RGBColor(0xC8, 0xC2, 0xB6)
CARD = RGBColor(0x1A, 0x1A, 0x1A)

PAGE_W = Inches(8.5)
PAGE_H = Inches(11)
MARGIN = Inches(0.5)

CONCEPTS = [
    {
        "id": "01",
        "name": "Subterranean Red",
        "role": "Simple SKU",
        "status": "Ready for review",
        "status_key": "ready",
        "pitch": "A full cherry field with the required 60 Years lockup front and center. The safest, most operational option — it reads on a concourse in one second.",
        "lockup": "Front panel, centered",
        "rights": "None. Graphic only.",
        "colors": "Cherry PMS 186 · White · Silver 877",
        "notes": [
            "Best as the default concession cup if only one SKU is produced.",
            "Correct the season line to 2026–27 before any print proof.",
            "Does not tell a unique Pit story beyond the required mark.",
        ],
    },
    {
        "id": "02",
        "name": "Remember the Feeling",
        "role": "Legacy / history",
        "status": "Ready for review",
        "status_key": "ready",
        "pitch": "Bob King’s line — “When you came down the ramp, it’s quite a feeling” — on a cherry-duotone panorama of the descent: the rail, the look down into the bowl, then the official 60 coin.",
        "lockup": "Front coin on the photo, ~2 in.  ·  Pit throwback mark on the black panel",
        "rights": "Yes. Historic / game photography must come from UNM Athletics archives.",
        "colors": "Cherry title band · Black · White  ·  official marks as supplied",
        "notes": [
            "The ramp is the story. Do not recrop this photo into a tight crowd shot.",
            "Title band must stay off the picture so REMEMBER THE FEELING reads on the cup.",
            "Pairs with The Opening as a two-SKU legacy set once photo rights are cleared.",
        ],
    },
    {
        "id": "03",
        "name": "Cherry Voltage",
        "role": "Modern merch",
        "status": "Studio recommendation",
        "status_key": "hero",
        "pitch": "Matte black. One oversized 60. A single cherry strike. The required lockup sits on the back as an untouched badge so the front can stay loud.",
        "lockup": "Back badge, 2 in., PNG as-is",
        "rights": "None. Graphic only.",
        "colors": "Matte black · White · Cherry strike only · Silver rules",
        "notes": [
            "Highest shelf impact of the set. Reads at 20 feet.",
            "Four-color / spot-friendly. No photography to clear.",
            "Do not let the cherry strike cut the official lockup.",
        ],
    },
    {
        "id": "04",
        "name": "Four Eras",
        "role": "Years in review",
        "status": "Concept hold — art in revision",
        "status_key": "hold",
        "pitch": "Turn the cup, move through 60 years: 1970s–80s cartoon wolf, 1990s bold color, early-2000s shield, 2026 lockup. The idea is sound. The drawings are not finished.",
        "lockup": "Panel 4 of 4 only (2026)",
        "rights": "Mascot art needs UNM Athletics approval. No licensed cartoon characters.",
        "colors": "Cream · Navy · Cherry · Teal (1990s only) · Black · White",
        "notes": [
            "Keep the four-panel structure. Do not approve current mascot drawings.",
            "1990s teal / LOBOS panel is the strongest piece of the wrap today.",
            "Present as a direction, not a print-ready face.",
        ],
    },
    {
        "id": "05",
        "name": "The Ramp",
        "role": "Architecture",
        "status": "Ready for review",
        "status_key": "ready",
        "pitch": "Type-led cream cup about the walk into the building. The Pit is underground. No other arena souvenir can own that sentence.",
        "lockup": "Base zone, below the type band",
        "rights": "None. Graphic only.",
        "colors": "Cream stock · Black type · Cherry rule · Official lockup",
        "notes": [
            "Quiet on purpose. Works as the “collector” or hospitality cup.",
            "Best story line for a buyer presentation: only we have the ramp.",
            "Keep it type-only. Do not add a photo dump later.",
        ],
    },
    {
        "id": "06",
        "name": "60 Howls",
        "role": "Collectible graphic",
        "status": "Concept hold — art in revision",
        "status_key": "hold",
        "pitch": "Sixty vertical bars — one per year — should form a howling wolf you see at 15 feet, then count at 2 feet. Current art reads as a red barcode cloud, not a wolf.",
        "lockup": "Side stamp (intended)",
        "rights": "None if the wolf is original geometry.",
        "colors": "Cream · Cherry · Black (3-color screenprint)",
        "notes": [
            "Retain the idea. Do not show this face as a finished option.",
            "Next pass needs a clear wolf-head silhouette filled with 60 bars.",
            "If the wolf never reads, retire the name and keep Cherry Voltage.",
        ],
    },
    {
        "id": "07",
        "name": "The Opening",
        "role": "Opening night / archive",
        "status": "Ready for review",
        "status_key": "ready",
        "pitch": "December 1, 1966. Lobos 62, Abilene Christian 53. A black-and-cherry duotone of the packed bowl, built like a Tribune front page.",
        "lockup": "Side commemorative stamp",
        "rights": "Yes. Period photography and any newspaper masthead need clearance.",
        "colors": "Black · Cherry duotone · White type",
        "notes": [
            "Newest concept on the board, and one of the strongest.",
            "Label every archive photo PHOTO · RIGHTS REQUIRED until legal clears it.",
            "Natural pair with Remember the Feeling for a history-led season.",
        ],
    },
]


def set_run(run, size, color, bold=False, name="Arial"):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = rpr.makeelement(qn("a:latin"), {})
        rpr.append(latin)
    latin.set("typeface", name)


def add_textbox(slide, l, t, w, h, text, size=11, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name="Arial"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.paragraphs[0].alignment = align
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size, color, bold, name)
    return box


def add_rect(slide, l, t, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def footer(slide, page, total):
    add_rect(slide, 0, Inches(10.72), PAGE_W, Inches(0.28), INK)
    add_rect(slide, 0, Inches(10.72), Inches(0.18), Inches(0.28), CHERRY)
    add_textbox(slide, MARGIN, Inches(10.74), Inches(5.4), Inches(0.22),
                "CONFIDENTIAL  ·  PROTOTYPE — NOT LICENSED FOR PRODUCTION  ·  UNM ATHLETICS",
                7, MUTED, False)
    add_textbox(slide, Inches(6.2), Inches(10.74), Inches(1.8), Inches(0.22),
                f"{page}  /  {total}", 7, MUTED, False, PP_ALIGN.RIGHT)


def chip(slide, l, t, w, h, text, key):
    fill = { "hero": CHERRY, "hold": RGBColor(0x3A, 0x2A, 0x14), "ready": RGBColor(0x1E, 0x2A, 0x22) }[key]
    ink = { "hero": WHITE, "hold": RGBColor(0xE6, 0xC0, 0x7B), "ready": RGBColor(0x9F, 0xC9, 0xA8) }[key]
    add_rect(slide, l, t, w, h, fill)
    add_textbox(slide, l + Inches(0.06), t + Inches(0.02), w - Inches(0.1), h - Inches(0.02),
                text.upper(), 8, ink, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def blank_black(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, PAGE_W, PAGE_H, INK)
    return slide


def cover(prs, total):
    slide = blank_black(prs)
    add_rect(slide, 0, 0, PAGE_W, Inches(0.12), CHERRY)
    slide.shapes.add_picture(str(ASSETS / "lockup-60.png"), Inches(2.85), Inches(1.15), Inches(2.8), Inches(2.8))
    add_textbox(slide, MARGIN, Inches(4.2), Inches(7.5), Inches(0.3),
                "UNM ATHLETICS  ·  THE PIT  ·  POWERED BY NUSENDA", 11, CHERRY, True, PP_ALIGN.CENTER)
    add_textbox(slide, MARGIN, Inches(4.6), Inches(7.5), Inches(1.4),
                "32 OZ SOUVENIR CUP", 36, WHITE, True, PP_ALIGN.CENTER)
    add_textbox(slide, MARGIN, Inches(5.95), Inches(7.5), Inches(0.4),
                "60th Anniversary  ·  Concept Review", 18, CREAM, False, PP_ALIGN.CENTER)
    add_rect(slide, Inches(3.7), Inches(6.5), Inches(1.1), Inches(0.04), CHERRY)
    add_textbox(slide, Inches(1.0), Inches(6.8), Inches(6.5), Inches(1.1),
                "Vertical print handout for executive review.\nSeven prototype faces. One required lockup. Not a production file.",
                13, WARM, False, PP_ALIGN.CENTER)
    add_textbox(slide, MARGIN, Inches(9.6), Inches(7.5), Inches(0.7),
                "SEPTEMBER 2026\nPACKAGING STUDY  ·  INTERNAL USE ONLY", 10, MUTED, False, PP_ALIGN.CENTER)
    footer(slide, 1, total)


def brief(prs, page, total):
    slide = blank_black(prs)
    add_rect(slide, 0, 0, PAGE_W, Inches(0.12), CHERRY)
    add_textbox(slide, MARGIN, Inches(0.35), Inches(7.5), Inches(0.25),
                "01  —  THE ASSIGNMENT", 10, CHERRY, True)
    add_textbox(slide, MARGIN, Inches(0.62), Inches(7.5), Inches(0.7),
                "What leadership is being asked to review", 24, WHITE, True)

    add_textbox(slide, MARGIN, Inches(1.5), Inches(7.5), Inches(1.15),
                "Design a 32 oz arena souvenir cup for the 60th anniversary season of University Arena — The Pit (1966–2026). Administration supplied one required mark and no other art direction. This packet translates the live prototype into a printed review.",
                13, CREAM)

    cards = [
        ("VESSEL", "32 fl oz plastic stadium cup\n~7 in. tall, tapered\nFull-wrap print, 10.5 × 5.75 in.\n0.25 in. bleed  ·  0.4 in. safe zone"),
        ("REQUIRED MARK", "Official 60 Years lockup, unaltered.\n1966 / 2026  ·  The Pit\nPowered by Nusenda\nLobo shield  ·  Bob King signature"),
        ("THIS PACKET", "One page per concept: cup, wrap,\nlockup placement, rights, and a\nstudio recommendation.\nPrototypes. Not licensed for production."),
    ]
    x = MARGIN
    for title, body in cards:
        add_rect(slide, x, Inches(2.85), Inches(2.4), Inches(2.35), CARD)
        add_rect(slide, x, Inches(2.85), Inches(2.4), Inches(0.08), CHERRY)
        add_textbox(slide, x + Inches(0.15), Inches(3.05), Inches(2.1), Inches(0.3), title, 10, CHERRY, True)
        add_textbox(slide, x + Inches(0.15), Inches(3.4), Inches(2.1), Inches(1.6), body, 11, CREAM)
        x += Inches(2.5)

    add_textbox(slide, MARGIN, Inches(5.45), Inches(7.5), Inches(0.3), "HOW TO READ THE SET", 10, CHERRY, True)
    bullets = [
        "Cherry Voltage is the studio’s modern recommendation.",
        "Remember the Feeling and The Opening are the history pair. Both need photo rights.",
        "The Ramp is the architecture story — unique to this building.",
        "Subterranean Red is the operational fallback: lockup on cherry.",
        "Four Eras and 60 Howls keep their briefs. Their current faces are not ready to approve.",
    ]
    y = Inches(5.85)
    for b in bullets:
        add_rect(slide, MARGIN, y + Inches(0.08), Inches(0.08), Inches(0.08), CHERRY)
        add_textbox(slide, Inches(0.75), y, Inches(7.2), Inches(0.4), b, 13, CREAM)
        y += Inches(0.48)

    add_textbox(slide, MARGIN, Inches(8.5), Inches(7.5), Inches(1.8),
                "Interactive source of these faces:\nhttps://stand-revel-66898789.figma.site",
                11, SILVER)
    footer(slide, page, total)


def lineup(prs, page, total):
    slide = blank_black(prs)
    add_rect(slide, 0, 0, PAGE_W, Inches(0.12), CHERRY)
    add_textbox(slide, MARGIN, Inches(0.35), Inches(7.5), Inches(0.25),
                "02  —  THE LINEUP", 10, CHERRY, True)
    add_textbox(slide, MARGIN, Inches(0.62), Inches(7.5), Inches(0.55),
                "Seven faces. One cup.", 24, WHITE, True)
    add_textbox(slide, MARGIN, Inches(1.2), Inches(7.5), Inches(0.5),
                "Printed from the live prototype. Hold Four Eras and 60 Howls at concept — do not treat those two drawings as finished art.",
                12, WARM)

    slide.shapes.add_picture(str(ASSETS / "00-lineup.png"), Inches(0.35), Inches(1.9), Inches(7.8), Inches(1.23))

    rows = [
        ("01", "Subterranean Red", "Simple SKU", "Ready"),
        ("02", "Remember the Feeling", "Legacy", "Ready"),
        ("03", "Cherry Voltage", "Modern", "Recommend"),
        ("04", "Four Eras", "Years in review", "Hold"),
        ("05", "The Ramp", "Architecture", "Ready"),
        ("06", "60 Howls", "Collectible", "Hold"),
        ("07", "The Opening", "Opening night", "Ready"),
    ]
    y = Inches(3.35)
    add_textbox(slide, MARGIN, y, Inches(0.6), Inches(0.28), "#", 9, MUTED, True)
    add_textbox(slide, Inches(1.2), y, Inches(3.2), Inches(0.28), "CONCEPT", 9, MUTED, True)
    add_textbox(slide, Inches(4.5), y, Inches(2.0), Inches(0.28), "JOB", 9, MUTED, True)
    add_textbox(slide, Inches(6.5), y, Inches(1.5), Inches(0.28), "STATUS", 9, MUTED, True)
    y += Inches(0.32)
    add_rect(slide, MARGIN, y, Inches(7.5), Inches(0.015), RGBColor(0x2A, 0x2A, 0x2A))
    y += Inches(0.12)
    for num, name, job, status in rows:
        add_textbox(slide, MARGIN, y, Inches(0.6), Inches(0.38), num, 14, CHERRY, True)
        add_textbox(slide, Inches(1.2), y, Inches(3.2), Inches(0.38), name, 14, WHITE, True)
        add_textbox(slide, Inches(4.5), y, Inches(2.0), Inches(0.38), job, 12, WARM)
        color = CHERRY if status == "Recommend" else (RGBColor(0xE6, 0xC0, 0x7B) if status == "Hold" else RGBColor(0x9F, 0xC9, 0xA8))
        add_textbox(slide, Inches(6.5), y, Inches(1.5), Inches(0.38), status.upper(), 11, color, True)
        y += Inches(0.52)

    add_textbox(slide, MARGIN, Inches(8.3), Inches(7.5), Inches(2.0),
                "Ask in the room: which one face should carry the season, and whether a second history SKU is worth the photo-rights work.",
                13, CREAM)
    footer(slide, page, total)


def lockup_page(prs, page, total):
    slide = blank_black(prs)
    add_rect(slide, 0, 0, PAGE_W, Inches(0.12), CHERRY)
    add_textbox(slide, MARGIN, Inches(0.35), Inches(7.5), Inches(0.25),
                "03  —  THE REQUIRED MARK", 10, CHERRY, True)
    add_textbox(slide, MARGIN, Inches(0.62), Inches(7.5), Inches(0.7),
                "This lockup does not get redesigned.", 22, WHITE, True)
    add_textbox(slide, MARGIN, Inches(1.4), Inches(7.5), Inches(0.9),
                "Cherry circle. Giant 60. YEARS ribbon. 1966 / 2026. THE PIT. POWERED BY NUSENDA. Lobo shield. Bob King signature. Place the official PNG. Never redraw, recolor, stretch, or drop a line.",
                13, CREAM)

    slide.shapes.add_picture(str(ASSETS / "lockup-60.png"), Inches(2.55), Inches(2.5), Inches(3.4), Inches(3.4))

    rules = [
        "Minimum width on cup: 1.75 in.  ·  Keep 0.4 in. off seam, rim, and base.",
        "Nusenda lives inside the mark. Do not add a second sponsor lockup.",
        "If a decorative line (the Cherry Voltage strike) approaches the badge, stop the line.",
        "Any generated version of this mark is a reject. Swap in the official file.",
    ]
    y = Inches(6.15)
    for r in rules:
        add_rect(slide, MARGIN, y + Inches(0.08), Inches(0.08), Inches(0.08), CHERRY)
        add_textbox(slide, Inches(0.75), y, Inches(7.2), Inches(0.45), r, 12, CREAM)
        y += Inches(0.5)
    footer(slide, page, total)


def concept_page(prs, concept, page, total):
    slide = blank_black(prs)
    add_rect(slide, 0, 0, PAGE_W, Inches(0.12), CHERRY)
    add_textbox(slide, MARGIN, Inches(0.28), Inches(4.6), Inches(0.22),
                f"CONCEPT  {concept['id']}  /  07   ·   {concept['role'].upper()}", 9, CHERRY, True)
    chip(slide, Inches(5.15), Inches(0.28), Inches(2.85), Inches(0.28), concept["status"], concept["status_key"])

    add_textbox(slide, MARGIN, Inches(0.62), Inches(7.5), Inches(0.55),
                concept["name"].upper(), 26, WHITE, True)
    add_textbox(slide, MARGIN, Inches(1.18), Inches(7.5), Inches(0.85),
                concept["pitch"], 12, CREAM)

    cup = ASSETS / f"{concept['id']}-cup.png"
    wrap = ASSETS / f"{concept['id']}-wrap.png"
    slide.shapes.add_picture(str(cup), Inches(0.4), Inches(2.15), Inches(3.05), Inches(4.44))

    x = Inches(3.65)
    y = Inches(2.15)
    meta = [
        ("REQUIRED MARK", concept["lockup"]),
        ("PRINT / COLOR", concept["colors"]),
        ("PHOTO RIGHTS", concept["rights"]),
    ]
    for label, body in meta:
        add_textbox(slide, x, y, Inches(4.35), Inches(0.22), label, 9, CHERRY, True)
        add_textbox(slide, x, y + Inches(0.22), Inches(4.35), Inches(0.55), body, 12, WHITE)
        y += Inches(0.82)

    add_textbox(slide, x, y, Inches(4.35), Inches(0.22), "STUDIO NOTES", 9, CHERRY, True)
    y += Inches(0.22)
    for note in concept["notes"]:
        add_textbox(slide, x, y, Inches(4.35), Inches(0.42), "·  " + note, 11, WARM)
        y += Inches(0.40)

    add_textbox(slide, MARGIN, Inches(6.78), Inches(7.5), Inches(0.2),
                "FULL WRAP  ·  10.5 × 5.75 IN.  ·  BACK  /  FRONT  /  SIDE", 9, CHERRY, True)
    slide.shapes.add_picture(str(wrap), Inches(0.5), Inches(7.02), Inches(7.5), Inches(3.5))
    footer(slide, page, total)


def recommend(prs, page, total):
    slide = blank_black(prs)
    add_rect(slide, 0, 0, PAGE_W, Inches(0.12), CHERRY)
    add_textbox(slide, MARGIN, Inches(0.35), Inches(7.5), Inches(0.25),
                "11  —  STUDIO RECOMMENDATION", 10, CHERRY, True)
    add_textbox(slide, MARGIN, Inches(0.62), Inches(7.5), Inches(0.7),
                "If one cup ships, ship Cherry Voltage.", 22, WHITE, True)
    add_textbox(slide, MARGIN, Inches(1.4), Inches(7.5), Inches(1.0),
                "It is the only face that already looks like a finished souvenir: high contrast, no photo rights, required mark legally parked on the back, cherry used as a strike instead of a wallpaper.",
                13, CREAM)

    blocks = [
        ("IF TWO SKUS", "Cherry Voltage for the concourse.\nRemember the Feeling or The Opening\nfor the history / alumni cup — after\nphoto rights are cleared."),
        ("IF YOU NEED SAFE", "Subterranean Red. Logo on cherry.\nIt will not win a design award.\nIt will not get anyone in trouble."),
        ("DO NOT APPROVE YET", "Four Eras and 60 Howls.\nKeep the briefs. Redo the drawings.\nShowing those faces as finals will\nweaken the rest of the room."),
    ]
    x = MARGIN
    for title, body in blocks:
        add_rect(slide, x, Inches(2.65), Inches(2.4), Inches(2.7), CARD)
        add_rect(slide, x, Inches(2.65), Inches(2.4), Inches(0.08), CHERRY)
        add_textbox(slide, x + Inches(0.14), Inches(2.85), Inches(2.12), Inches(0.55), title, 11, CHERRY, True)
        add_textbox(slide, x + Inches(0.14), Inches(3.4), Inches(2.12), Inches(1.75), body, 12, CREAM)
        x += Inches(2.5)

    slide.shapes.add_picture(str(ASSETS / "03-cup.png"), Inches(2.7), Inches(5.6), Inches(3.1), Inches(4.5))
    add_textbox(slide, MARGIN, Inches(5.6), Inches(2.1), Inches(1.4),
                "CHERRY\nVOLTAGE\n\n03 / 07", 14, WHITE, True)
    footer(slide, page, total)


def next_steps(prs, page, total):
    slide = blank_black(prs)
    add_rect(slide, 0, 0, PAGE_W, Inches(0.12), CHERRY)
    add_textbox(slide, MARGIN, Inches(0.35), Inches(7.5), Inches(0.25),
                "12  —  NEXT STEPS  ·  FINE PRINT", 10, CHERRY, True)
    add_textbox(slide, MARGIN, Inches(0.62), Inches(7.5), Inches(0.55),
                "What to decide in this meeting", 22, WHITE, True)

    steps = [
        ("1", "Pick a primary face for the 2026–27 souvenir cup."),
        ("2", "Decide whether a second history SKU is in budget (photo rights + a second plate)."),
        ("3", "Confirm the official 60 Years PNG is the only lockup that may print."),
        ("4", "Park Four Eras and 60 Howls as R&D unless new art is requested."),
        ("5", "Move the approved wrap to a real 32 oz stadium-cup dieline — current mockups are still a pint silhouette."),
    ]
    y = Inches(1.4)
    for n, text in steps:
        add_rect(slide, MARGIN, y, Inches(0.38), Inches(0.38), CHERRY)
        add_textbox(slide, MARGIN, y + Inches(0.04), Inches(0.38), Inches(0.32), n, 14, WHITE, True, PP_ALIGN.CENTER)
        add_textbox(slide, Inches(1.05), y + Inches(0.02), Inches(6.9), Inches(0.5), text, 14, CREAM)
        y += Inches(0.62)

    add_rect(slide, MARGIN, Inches(5.0), Inches(7.5), Inches(0.015), RGBColor(0x2A, 0x2A, 0x2A))
    add_textbox(slide, MARGIN, Inches(5.2), Inches(7.5), Inches(0.3), "FINE PRINT", 10, CHERRY, True)
    add_textbox(slide, MARGIN, Inches(5.55), Inches(7.5), Inches(3.4),
                "These are prototype wraps for internal review. They are not licensed for production, sale, or public release. University marks, The Pit naming, Nusenda partnership language, and Bob King’s signature are used here only as specified by the official 60th lockup. Historic photography and any newspaper treatment remain placeholders until UNM Athletics and Legal clear them.\n\nVessel note: renderings use a pint-style silhouette. Production stock should be a 32 oz stadium cup.\n\nInteractive reference: https://stand-revel-66898789.figma.site\n\nPrepared as a printed executive handout, September 2026.",
                12, WARM)
    slide.shapes.add_picture(str(ASSETS / "lockup-60.png"), Inches(6.15), Inches(8.55), Inches(1.7), Inches(1.7))
    footer(slide, page, total)


def build():
    prs = Presentation()
    prs.slide_width = PAGE_W
    prs.slide_height = PAGE_H
    total = 4 + len(CONCEPTS) + 2

    cover(prs, total)
    brief(prs, 2, total)
    lineup(prs, 3, total)
    lockup_page(prs, 4, total)
    page = 5
    for concept in CONCEPTS:
        concept_page(prs, concept, page, total)
        page += 1
    recommend(prs, page, total)
    next_steps(prs, page + 1, total)

    prs.save(OUT)
    print(f"Wrote {OUT} ({total} slides)")


if __name__ == "__main__":
    build()
